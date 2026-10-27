"""Build the one-page project abstract: ../abstract.docx, then ../abstract.pdf.

    python build_abstract.py

Layout: A4, Times New Roman, a bold 16 pt centred title, a centred "Abstract" heading, justified 12 pt body
text, the student's name and roll number in bold towards the right, and the guide's signature block below.
Every figure in the text is worked out, with its source, in ../benchmark-delta-x.md.
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from build_references import page_count

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "abstract.docx")
FONT = "Times New Roman"

TITLE = "Predictive Thermal Battery Management System for Li-ion Battery Packs"
STUDENT = "Ms. Muskan Sulathana (23WH1A0208)"
SIGNATURE = ["Signature", "Dr. M. Rupesh", "Associate Professor", "EEE Department"]

BODY = [
    "An internal short inside one cell is often named as the cause of lithium-ion battery accidents. In its "
    "most common form it starts small: it slowly drains its own cell and makes only a fraction of a watt of "
    "heat. A normal battery management system (BMS) cannot see this stage, because it acts only when a "
    "voltage, current or temperature crosses a fixed limit. In this project I am building a four-cell (4S1P) "
    "NCR18650GA battery pack with two layers of protection: an independent hardware BMS, and an ESP32 "
    "supervisor that predicts. Five times a second, the ESP32 runs a physical thermal model of every cell "
    "with a Kalman filter, and estimates what no sensor can measure: the core temperature, and the abnormal "
    "heat, in watts, that the cell's own current cannot explain. It forecasts each cell's temperature 1 and "
    "5 minutes ahead and checks every forecast against the real reading, and at rest it estimates each "
    "cell's self-discharge from its voltage. All of this evidence becomes one risk level (Normal, Watch, "
    "Warning or Trip), which can open separate charge and load relays and keep them latched.",

    "I benchmarked the design against six commercial products, from a ₹239 hardware BMS board and protection "
    "chips from TI, ADI and NXP to the Orion BMS and the Li-ion Tamer gas sensor. All of them handle "
    "temperature with fixed limits, and the only early warning among them, the gas sensor (US$596), reacts "
    "only after a cell has vented. Among seven published methods, the longest warnings come from methods "
    "trained on cells or a vehicle taken into thermal runaway, and only one of them acts on the battery.",

    "The Delta X over the base paper, Ojo et al. (IEEE TIE, 2021), which flags faults with a trained neural "
    "network, is as follows. My design needs no training data: 11 physical constants, fitted once from "
    "normal use, replace about 21,000 trained weights. The base paper's 0.69 °C threshold within its 4.5 s "
    "window catches heating only above about 9 °C/min, while my design is built to warn at 0.5 °C/min "
    "(about 0.4 W), 18 times slower, and its 5-minute check catches 0.2 W. It adds electrical evidence of a "
    "short, a leak from 20 mA; it needs 10–30 s of persistence and agreement between thermal and electrical "
    "evidence before an alarm; it uses about 2,000 times less arithmetic per estimate on a ₹639 ESP32 "
    "instead of a Jetson TX2 or Raspberry Pi; and it disconnects the pack instead of only raising a flag.",

    "Semester 1 builds and calibrates the electronics and the first predictive firmware with no lithium "
    "cells (Gates 0–2). Semester 2 builds the pack and tests every target at Gates 9–11, including a safely "
    "emulated soft short. Until those tests, all figures for my design are targets, not results.",
]


def keep_units(text):
    """Join each number to its unit with a non-breaking space, so a value never splits across lines."""
    return re.sub(r"(\d) (°C|s\b|W\b|mA\b|minutes\b)", "\\1\u00a0\\2", text)


def paragraph(doc, text="", size=12, bold=False, align=None, before=0, after=0, line=1.08,
              left=None, right=None):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Twips(before), Twips(after), line
    if align is not None:
        p.alignment = align
    if left is not None:
        f.left_indent = Twips(left)
    if right is not None:
        f.right_indent = Twips(right)
    if text:
        run = p.add_run(keep_units(text))
        run.bold, run.font.size = bold, Pt(size)
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11910), Twips(16840)             # A4
    sec.top_margin, sec.bottom_margin = Twips(1360), Twips(720)
    sec.left_margin = sec.right_margin = Twips(1417)                         # 2.5 cm

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    zoom = doc.settings.element.find(qn("w:zoom"))      # the default template omits the required percent
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")

    core = doc.core_properties
    core.author = core.last_modified_by = "Muskan Sulathana"
    core.title, core.subject, core.comments, core.keywords = TITLE, "Project abstract", "", ""

    paragraph(doc, TITLE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=350, after=120)
    paragraph(doc, "Abstract", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=240)
    for text in BODY:
        paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=120)

    paragraph(doc, STUDENT, bold=True, before=360, left=4344, line=398 / 240)
    for i, line in enumerate(SIGNATURE):
        paragraph(doc, line, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  before=420 if i == 0 else 0, left=5600, line=379 / 240)

    doc.save(OUT)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(OUT), OUT],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pages = page_count(os.path.splitext(OUT)[0] + ".pdf")
    print(f"saved abstract.docx and abstract.pdf ({pages} page{'s' if pages != 1 else ''})")
    if pages != 1:
        sys.exit("the abstract must fit on one page; shorten BODY")


if __name__ == "__main__":
    build()
