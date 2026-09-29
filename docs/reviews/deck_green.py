"""Slide furniture for the green Review-1 template.

Measured from the department's Review-1 decks (see template-notes.md): 4:3
pages, a green bar across the top and a bevelled green bar at the foot, plain
left-aligned headings in Times New Roman, and blue Office-style tables.

The low-level helpers (text boxes, runs, bullets, overflow check, notes) come
from deck_common.py, so both templates write text the same way. A build script
calls configure() first, then uses the slide builders below.
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from deck_common import (textbox, run, set_bullet, no_bullet, notes, _fit,
                         LOGO, BODY_FONT, HEADING_FONT, BULLET_FONT)

# ── page and colour tokens, measured from the template ──────────────────
SLIDE_W, SLIDE_H = 10.0, 7.5

GREEN = RGBColor(0x4A, 0x7A, 0x30)
EDGE = RGBColor(0x16, 0x16, 0x16)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY_TEXT = RGBColor(0x40, 0x40, 0x40)

TITLE_NAVY = RGBColor(0x00, 0x20, 0x60)      # "Major Project Stage-1 ..."
TITLE_RED = RGBColor(0xFF, 0x00, 0x00)       # project title
GUIDE_BLUE = RGBColor(0x1F, 0x38, 0x64)      # supervisor lines
DEPT_PURPLE = RGBColor(0x80, 0x00, 0x80)     # department line
COLLEGE_NAVY = RGBColor(0x17, 0x36, 0x5D)    # college line

TABLE_HEAD = RGBColor(0x4F, 0x81, 0xBD)      # Office accent 1
TABLE_ROW = RGBColor(0xD0, 0xD8, 0xE8)

HEAD_L, HEAD_T, HEAD_H, HEAD_SIZE = 0.45, 0.52, 0.80, 40
BODY_L, BODY_R = 0.50, 9.56
BODY_T, BODY_B = 1.45, 6.85
BODY_SIZE = 20

# settings that differ between decks
EXPORT_DATE = ""
DIAG = None
REVIEW_HEADING = ""
TITLE_LINES = []
PRESENTER = ""
ROLL_NO = ""
SUPERVISOR_LINES = []

WARNINGS = []


def configure(export_date, diagram_dir, review_heading, title_lines,
              presenter, roll_no, supervisor_lines):
    global EXPORT_DATE, DIAG, REVIEW_HEADING, TITLE_LINES
    global PRESENTER, ROLL_NO, SUPERVISOR_LINES
    EXPORT_DATE = export_date
    DIAG = diagram_dir
    REVIEW_HEADING = review_heading
    TITLE_LINES = list(title_lines)
    PRESENTER = presenter
    ROLL_NO = roll_no
    SUPERVISOR_LINES = list(supervisor_lines)
    WARNINGS.clear()


def new_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def fits(name, lines, size, width_in, height_in, spacing=1.20, gap_pt=0):
    """Rough fit check for Times New Roman: python-pptx cannot shrink text."""
    chars_per_line = max(1, int(width_in * 72 / (size * 0.47)))
    n = sum(max(1, -(-len(s) // chars_per_line)) for s in lines)
    needed = n * size * spacing / 72.0 + max(0, len(lines) - 1) * gap_pt / 72.0
    if needed > height_in:
        WARNINGS.append("%s: needs ~%.2f in, box is %.2f in" % (name, needed, height_in))


# ── the green bars ──────────────────────────────────────────────────────
def _poly(slide, pts, color):
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=True)
    shape = fb.convert_to_shape()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _bars(slide, title=False):
    if title:
        _poly(slide, [(0.12, 0.20), (9.87, 0.20), (9.46, 0.37), (0.53, 0.37)], EDGE)
        _poly(slide, [(0.12, 0.00), (9.87, 0.00), (9.87, 0.23), (0.12, 0.23)], GREEN)
        _poly(slide, [(0.55, 6.93), (9.45, 6.93), (9.55, 7.02), (0.45, 7.02)], EDGE)
        _poly(slide, [(0.45, 7.02), (9.55, 7.02), (9.84, 7.31), (0.24, 7.31)], GREEN)
    else:
        _poly(slide, [(0.12, 0.24), (9.87, 0.24), (9.87, 0.29), (9.78, 0.37), (0.12, 0.37)], EDGE)
        _poly(slide, [(0.12, 0.00), (9.87, 0.00), (9.87, 0.27), (0.12, 0.27)], GREEN)
        _poly(slide, [(1.17, 7.02), (8.72, 7.02), (8.80, 7.10), (1.09, 7.10)], EDGE)
        _poly(slide, [(1.09, 7.10), (8.80, 7.10), (8.97, 7.31), (0.92, 7.31)], GREEN)


def base(prs, heading=None, number=None, heading_size=HEAD_SIZE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bars(slide)
    if heading:
        _, tf = textbox(slide, HEAD_L, HEAD_T, SLIDE_W - 2 * HEAD_L, HEAD_H)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run(tf.paragraphs[0], heading, heading_size, font=HEADING_FONT)
    if number is not None:
        _, tf = textbox(slide, 8.98, 7.08, 0.45, 0.22)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run(p, str(number), 10, color=GREY_TEXT)
    return slide


def _bullets(tf, items, size, gap_pt, justify=True, line=1.0):
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap_pt)
        p.line_spacing = line
        if justify:
            p.alignment = PP_ALIGN.JUSTIFY
        set_bullet(p, indent_in=0.24)
        if isinstance(text, tuple):                 # (bold lead-in, rest)
            run(p, text[0], size, bold=True)
            run(p, text[1], size)
        else:
            run(p, text, size)


# ── slide builders ──────────────────────────────────────────────────────
def title_slide(prs, note):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bars(slide, title=True)

    _, tf = textbox(slide, 7.60, 0.72, 1.95, 0.32)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run(p, EXPORT_DATE, 16)

    _, tf = textbox(slide, 0.60, 0.44, 8.80, 0.90)
    for i, text in enumerate([REVIEW_HEADING, "on"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run(p, text, 22, bold=True, color=TITLE_NAVY, font=HEADING_FONT)

    _, tf = textbox(slide, 0.45, 1.30, 9.10, 0.95)
    for i, text in enumerate(TITLE_LINES):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run(p, text, 26, bold=True, color=TITLE_RED, font=HEADING_FONT)

    logo_w = 1.00
    slide.shapes.add_picture(LOGO, Inches((SLIDE_W - logo_w) / 2), Inches(2.38),
                             Inches(logo_w), Inches(logo_w * 648 / 448))

    _, tf = textbox(slide, 1.35, 4.22, 4.3, 0.40)
    run(tf.paragraphs[0], "Presented by", 20, bold=True, color=GREY_TEXT)
    _, tf = textbox(slide, 1.38, 4.70, 2.60, 0.35)
    run(tf.paragraphs[0], PRESENTER, 18)
    _, tf = textbox(slide, 3.95, 4.70, 1.90, 0.35)
    run(tf.paragraphs[0], ":  " + ROLL_NO, 18)

    _, tf = textbox(slide, 6.05, 4.28, 3.10, 1.40)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, "Supervised by", 18, bold=True)
    for text in SUPERVISOR_LINES:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run(p, text, 18, color=GUIDE_BLUE)

    _, tf = textbox(slide, 0.30, 5.98, 9.40, 0.85)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, "Department of Electrical & Electronics Engineering", 22, bold=True,
        color=DEPT_PURPLE, font=HEADING_FONT)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run(p, "BVRIT HYDERABAD College of Engineering for Women", 22, bold=True,
        color=COLLEGE_NAVY, font=HEADING_FONT)

    notes(slide, note)
    return slide


def bullet_slide(prs, n, heading, items, note, size=BODY_SIZE, gap_pt=14,
                 justify=True, top=BODY_T, heading_size=HEAD_SIZE):
    slide = base(prs, heading, n, heading_size=heading_size)
    _, tf = textbox(slide, BODY_L, top, BODY_R - BODY_L, BODY_B - top)
    _bullets(tf, items, size, gap_pt, justify)
    plain = [i if isinstance(i, str) else i[0] + i[1] for i in items]
    fits(heading, plain, size, BODY_R - BODY_L - 0.24, BODY_B - top, gap_pt=gap_pt)
    notes(slide, note)
    return slide


def two_section_slide(prs, n, h1, items1, h2, items2, note, size=18, gap_pt=8):
    """Two headings on one page, as the template does Problem Analysis + Research Gap."""
    slide = base(prs, h1, n)
    split = 1.30 + 0.30 + (BODY_B - 1.30) * 0.50
    _, tf = textbox(slide, BODY_L, 1.30, BODY_R - BODY_L, split - 1.30 - 0.70)
    _bullets(tf, items1, size, gap_pt)
    plain1 = [i if isinstance(i, str) else "".join(i) for i in items1]
    fits(h1, plain1, size, BODY_R - BODY_L - 0.24, split - 2.00, gap_pt=gap_pt)

    _, tf = textbox(slide, HEAD_L, split - 0.62, SLIDE_W - 2 * HEAD_L, 0.72)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    run(tf.paragraphs[0], h2, HEAD_SIZE, font=HEADING_FONT)

    _, tf = textbox(slide, BODY_L, split + 0.16, BODY_R - BODY_L, BODY_B - split - 0.16)
    _bullets(tf, items2, size, gap_pt)
    plain2 = [i if isinstance(i, str) else "".join(i) for i in items2]
    fits(h2, plain2, size, BODY_R - BODY_L - 0.24, BODY_B - split - 0.16, gap_pt=gap_pt)
    notes(slide, note)
    return slide


def table_slide(prs, n, heading, headers, rows, widths, note, size=13,
                head_size=None, row_h=None, aligns=None, top=1.42,
                bold_first_col=True, legend=None, heading_size=HEAD_SIZE):
    """Blue Office-style table: accent-1 header, one flat body tint, white rules."""
    slide = base(prs, heading, n, heading_size=heading_size)
    total_w = sum(widths)
    left = (SLIDE_W - total_w) / 2
    nrows = len(rows) + 1
    avail = (6.30 if legend else BODY_B) - top
    height = min(avail, row_h * nrows) if row_h else avail
    shape = slide.shapes.add_table(nrows, len(headers), Inches(left), Inches(top),
                                   Inches(total_w), Inches(height))
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    # Pin the header to one short band and give the body rows the rest, the way
    # the template does, instead of spreading the height evenly over every row.
    head_h = min(0.50, height / nrows)
    table.rows[0].height = Inches(head_h)
    for r in range(1, nrows):
        table.rows[r].height = Inches((height - head_h) / (nrows - 1))
    aligns = aligns or (["c"] + ["l"] * (len(headers) - 1))

    def fill(cell, color, text, sz, bold, align):
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        tf = cell.text_frame
        tf.word_wrap = True
        parts = text if isinstance(text, (list, tuple)) else [text]
        for k, part in enumerate(parts):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER if align == "c" else PP_ALIGN.LEFT
            if isinstance(part, dict):                      # {"b": lead, "t": rest}
                run(p, part["b"], sz, bold=True)
                run(p, part["t"], sz)
            else:
                run(p, part, sz, bold=bold)

    for c, h in enumerate(headers):
        fill(table.cell(0, c), TABLE_HEAD, h, head_size or size + 1, True, "c")
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            fill(table.cell(r, c), TABLE_ROW, val, size,
                 bold_first_col and c == 0, aligns[c])

    if legend:
        _, tf = textbox(slide, 6.80, 6.36, 2.80, 0.62)
        for k, text in enumerate(legend):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.RIGHT
            run(p, text, 17)
    notes(slide, note)
    return slide


def image_slide(prs, n, heading, image, note, box=None, caption=None,
                heading_size=HEAD_SIZE):
    slide = base(prs, heading, n, heading_size=heading_size)
    l, t, w, h = box or (0.40, 1.40, 9.20, 5.50)
    if caption:
        h -= 0.36
    il, it, iw, ih = _fit(os.path.join(DIAG, image), l, t, w, h)
    slide.shapes.add_picture(os.path.join(DIAG, image), Inches(il), Inches(it),
                             Inches(iw), Inches(ih))
    if caption:
        _, tf = textbox(slide, l, t + h + 0.06, w, 0.30)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, caption, 13, italic=True, color=GREY_TEXT)
    notes(slide, note)
    return slide


def references_slide(prs, n, heading, refs, start, note, size=14, gap_pt=10):
    slide = base(prs, heading, n) if heading else base(prs, None, n)
    top = 1.40 if heading else 0.80
    _, tf = textbox(slide, BODY_L - 0.12, top, BODY_R - BODY_L + 0.12, BODY_B - top)
    for i, ref in enumerate(refs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap_pt)
        p.alignment = PP_ALIGN.LEFT            # justified spacing tears apart long DOIs
        no_bullet(p)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.46)).emu))
        pPr.set("indent", str(-Emu(Inches(0.46)).emu))
        run(p, "[%d]\t" % (start + i), size)
        run(p, ref, size)
    fits("References %d–%d" % (start, start + len(refs) - 1), refs, size,
         BODY_R - BODY_L - 0.4, BODY_B - top, gap_pt=gap_pt)
    notes(slide, note)
    return slide


def thank_you_slide(prs, n, note):
    slide = base(prs, None, n)
    _, tf = textbox(slide, 0.5, 2.85, 9.0, 1.4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, "THANK YOU", 66, font=HEADING_FONT)
    notes(slide, note)
    return slide
